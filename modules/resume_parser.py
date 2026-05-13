"""
modules/resume_parser.py  —  Phase 1 upgrade
Supports: PDF, DOCX, TXT, MD
"""

import os
import re
from dotenv import load_dotenv
import pdfplumber
from groq import Groq

load_dotenv()

GROQ_API_KEY = os.getenv('GROQ_API_KEY')
if not GROQ_API_KEY:
    raise ValueError("GROQ_API_KEY not set in environment.")

client = Groq(api_key=GROQ_API_KEY)
# llama3-70b-8192 was decommissioned Aug 2025 — replaced with successor
MODEL  = 'llama-3.3-70b-versatile'


# ── Text Extractors ───────────────────────────────────────────────────────────

def extract_text_from_pdf(pdf_path: str) -> str:
    text = ""
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text.strip()


def extract_text_from_docx(docx_path: str) -> str:
    """Extract text from .docx using python-docx."""
    try:
        from docx import Document
    except ImportError:
        raise ImportError("python-docx not installed. Run: pip install python-docx")

    doc   = Document(docx_path)
    lines = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = ' | '.join(
                cell.text.strip() for cell in row.cells if cell.text.strip()
            )
            if row_text:
                lines.append(row_text)

    return '\n'.join(lines).strip()


def extract_text_from_txt(path: str) -> str:
    with open(path, 'r', encoding='utf-8') as f:
        return f.read().strip()


def load_resume(file_path: str) -> str:
    """Load resume text from PDF, DOCX, TXT, or MD file."""
    ext = os.path.splitext(file_path)[-1].lower()

    if ext == '.pdf':
        return extract_text_from_pdf(file_path)
    elif ext == '.docx':
        return extract_text_from_docx(file_path)
    elif ext in ['.txt', '.md']:
        return extract_text_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}. Supported: .pdf, .docx, .txt, .md")


# ── Section Extractor (rule-based pre-processing) ─────────────────────────────

SECTION_HEADERS = {
    'skills':     r'(skills|technical\s+skills|core\s+competencies|technologies)',
    'experience': r'(experience|work\s+experience|employment|professional\s+experience)',
    'projects':   r'(projects|personal\s+projects|academic\s+projects|key\s+projects)',
    'education':  r'(education|academic|qualification)',
}

def _extract_sections(raw_text: str) -> dict:
    """
    Heuristically split resume text into sections.
    Returns dict with section names as keys and extracted text as values.
    """
    sections     = {k: '' for k in SECTION_HEADERS}
    lines        = raw_text.split('\n')
    current      = None
    buffer       = []

    for line in lines:
        stripped = line.strip()
        matched  = False
        for section, pattern in SECTION_HEADERS.items():
            if re.match(pattern, stripped, re.IGNORECASE) and len(stripped) < 60:
                # Save previous section
                if current:
                    sections[current] = '\n'.join(buffer).strip()
                current = section
                buffer  = []
                matched = True
                break
        if not matched and current:
            buffer.append(line)

    if current and buffer:
        sections[current] = '\n'.join(buffer).strip()

    return sections


# ── LLM-based Parser ──────────────────────────────────────────────────────────

def parse_resume(file_path: str) -> dict:
    """
    Full parse: load → extract sections → LLM structured extraction.
    Returns dict with name, job_role, experience, skills, projects, summary, raw_text.
    """
    raw_text = load_resume(file_path)

    if not raw_text or len(raw_text) < 50:
        raise ValueError("Resume appears empty or unreadable.")

    # Pre-extract sections to give LLM focused context
    sections       = _extract_sections(raw_text)
    skills_hint    = sections['skills'][:500]    if sections['skills']    else ''
    projects_hint  = sections['projects'][:800]  if sections['projects']  else ''
    experience_hint= sections['experience'][:800]if sections['experience']else ''

    prompt = f"""You are a precise resume parser. Extract information from the resume below.
Respond ONLY in this exact format — no extra text, no markdown, no preamble:

Name: <full name>
Job Role: <target or most recent job role>
Experience: <total years, e.g. "2 years" or "Fresher">
Skills: <comma-separated top 8-10 technical skills>
Projects: <pipe-separated list of project names, e.g. "Project A | Project B">
Education: <highest degree and institution>
Summary: <2-3 sentence professional summary focused on strengths>

Resume Text:
{raw_text[:3000]}

---
Extracted Sections for reference:
Skills Section: {skills_hint}
Projects Section: {projects_hint}
Experience Section: {experience_hint}"""

    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.2,
    )

    parsed_text = response.choices[0].message.content
    if not parsed_text:
        raise ValueError("Empty response from LLM during resume parsing.")

    parsed_text = parsed_text.strip()
    result      = {"raw_text": raw_text, "sections": sections}

    for line in parsed_text.split('\n'):
        if ':' not in line:
            continue
        key, _, value = line.partition(':')
        key   = key.strip().lower().replace(' ', '_')
        value = value.strip()
        if value:
            result[key] = value

    return result


# ── Public API ────────────────────────────────────────────────────────────────

def resume_to_profile(file_path: str) -> dict:
    """
    Parse resume and return a clean profile dict ready to feed into llm.py.
    """
    parsed = parse_resume(file_path)

    # Parse projects list
    projects_raw = parsed.get('projects', '')
    projects     = [p.strip() for p in projects_raw.split('|') if p.strip()]

    return {
        'name':        parsed.get('name',       'Candidate'),
        'job_role':    parsed.get('job_role',   'Software Engineer'),
        'experience':  parsed.get('experience', 'Fresher'),
        'skills':      parsed.get('skills',     ''),
        'projects':    projects,
        'education':   parsed.get('education',  ''),
        'summary':     parsed.get('summary',    ''),
        'resume_text': parsed.get('raw_text',   ''),
        'sections':    parsed.get('sections',   {}),
    }


def get_resume_context_for_llm(profile: dict) -> str:
    """
    Format profile into a concise context string to inject into LLM prompts.
    Use this when passing resume info to generate_questions() in llm.py.
    """
    lines = [
        f"Candidate: {profile.get('name', '')}",
        f"Target Role: {profile.get('job_role', '')}",
        f"Experience: {profile.get('experience', '')}",
        f"Skills: {profile.get('skills', '')}",
    ]
    if profile.get('projects'):
        lines.append(f"Notable Projects: {', '.join(profile['projects'][:3])}")
    if profile.get('education'):
        lines.append(f"Education: {profile.get('education', '')}")
    if profile.get('summary'):
        lines.append(f"Summary: {profile.get('summary', '')}")

    return '\n'.join(lines)